"""
Plain-Python (no Streamlit dependency) helpers for turning the files in
resources/ into text chunks ready for embedding.

This is intentionally separate from utils/qa.py so it can be imported by:
  - utils/qa.py itself, at app runtime
  - scripts/build_embeddings.py, an offline CLI script with no Streamlit
    process around it

Keeping the extraction/chunking logic in one place means the offline script
and the live app can never drift out of sync with each other.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

APP_DIR = Path(__file__).parent.parent
RESOURCES_DIR = APP_DIR / "resources"
MANIFEST_PATH = RESOURCES_DIR / "manifest.json"
EMBEDDINGS_PATH = RESOURCES_DIR / "embeddings.npz"

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 150


def manifest_hash(manifest_json: str) -> str:
    """Stable fingerprint of the manifest contents. Used to detect when the
    offline embeddings file is stale relative to resources/manifest.json."""
    return hashlib.sha256(manifest_json.encode("utf-8")).hexdigest()


def extract_text(file_path: Path) -> str:
    """Best-effort plain-text extraction for pptx, docx, pdf, txt, and md
    resources. Returns "" (rather than raising) for anything that fails to
    parse, so one bad file can't break indexing for the rest of the library.
    """
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(file_path))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)

        if suffix == ".docx":
            import docx

            doc = docx.Document(str(file_path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)

        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            return "\n".join((page.extract_text() or "") for page in reader.pages)

        if suffix in (".txt", ".md"):
            return file_path.read_text(encoding="utf-8")
    except Exception:
        return ""
    return ""


def chunk_text(text: str) -> list[str]:
    """Split into overlapping character chunks. Simple and good enough for
    slide-deck / Q&A-doc sized resources — no sentence-boundary awareness
    needed for this scale of content."""
    text = " ".join(text.split())
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start = end - CHUNK_OVERLAP
    return chunks


def iter_source_chunks(manifest: list[dict]) -> list[tuple[str, str, str]]:
    """Given a parsed manifest, return (chunk_text, source_title,
    source_topic) tuples for every chunk of every local resource file.
    External `url` entries are skipped — there's no local text to read."""
    results: list[tuple[str, str, str]] = []
    for r in manifest:
        if "file" not in r:
            continue
        file_path = RESOURCES_DIR / r["file"]
        if not file_path.exists():
            continue
        text = extract_text(file_path)
        for piece in chunk_text(text):
            results.append((piece, r["title"], r.get("topic", "")))
    return results


def load_manifest() -> list[dict]:
    return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
